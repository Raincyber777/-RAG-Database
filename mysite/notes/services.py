"""
纪要生成流水线（真实版）：上传材料 → 按类型抽取内容 → LLM 融合生成纪要板块 → 向量化入库
- audio/video：DashScope paraformer-v2 文件转写（本地文件上传临时 OSS → 异步任务轮询 →
  句子级时间戳），转写行带 [mm:ss] 前缀，供 RAG 引用对齐到录音时间点
- image：qwen-vl-plus 识别板书 / PPT 截图
- doc / ppt：txt / md 直读；pdf（PyMuPDF）、docx（python-docx）、pptx / ppsx（python-pptx）、
  xlsx（openpyxl）、csv 按扩展名分派解析；老版二进制 .ppt / .doc / .xls 在序列化层即拒绝
- 抽取失败 / LLM 失败：状态置为 failed，接口层不阻塞（全程后台线程）
"""
import base64
import logging
import mimetypes
import os
import re
import threading
import time
import uuid

import requests
from django.conf import settings

log = logging.getLogger(__name__)

from mysite import llm, rag
from mysite.media_storage import local_material_path
from notes.models import Note

ASR_BASE = 'https://dashscope.aliyuncs.com'
ASR_MODEL = 'paraformer-v2'
ASR_POLL_INTERVAL = 3      # 秒
ASR_POLL_MAX = 400         # 最长轮询 20 分钟（一节课录音量级）
NO_SPEECH_TEXT = ('未在材料中识别到可用的语音内容：请确认音视频里有人声讲解'
                  '（纯音乐 / 纯画面 / 静音材料无法生成纪要），可更换材料后重试。')
TEXT_MAX = 60000           # 文本材料保留的最大字符数（RAG 入库与转写原文共用）
SUMMARIZE_DIRECT_MAX = 15000   # 不超过该长度直接单次总结；超出走分段 map-reduce
SUMMARIZE_CHUNK = 12000        # map-reduce 每段目标长度（按段落边界切）
ZIP_INNER_MAX = 20 * 1024 * 1024    # 压缩包内单文件解压上限（对齐图片上限量级）
ZIP_TOTAL_MAX = 200 * 1024 * 1024   # 压缩包解压后总量上限（防压缩炸弹）


class UnsupportedSource(Exception):
    """材料类型暂不支持自动生成"""


def _mmss(milliseconds: int) -> str:
    total = max(0, int(milliseconds)) // 1000
    return f'{total // 60:02d}:{total % 60:02d}'


def _upload_to_oss(file_path: str) -> str:
    """本地文件上传到 DashScope 临时 OSS，返回 oss:// 引用（录音文件识别要求可访问 URL）"""
    filename = f'{uuid.uuid4().hex}-{os.path.basename(file_path)}'
    res = requests.get(
        f'{ASR_BASE}/api/v1/uploads',
        params={'action': 'getPolicy', 'model': ASR_MODEL},
        headers={'Authorization': f'Bearer {settings.DASHSCOPE_API_KEY}'},
        timeout=30,
    )
    res.raise_for_status()
    pol = res.json()['data']
    # upload_dir 不带尾部斜杠（形如 dashscope-instant/<uid>/<date>/<task-uuid>），需显式补 /
    obj_key = pol['upload_dir'].rstrip('/') + '/' + filename
    form = {
        'OSSAccessKeyId': pol['oss_access_key_id'],
        'Signature': pol['signature'],
        'policy': pol['policy'],
        'x-oss-object-acl': pol.get('x_oss_object_acl', 'private'),
        'x-oss-forbid-overwrite': pol.get('x_oss_forbid_overwrite', 'true'),
        'key': obj_key,
        'success_action_status': '200',
    }
    with open(file_path, 'rb') as f:
        res = requests.post(pol['upload_host'], data=form, files={'file': f}, timeout=300)
    res.raise_for_status()
    return 'oss://' + obj_key


class NoSpeechError(RuntimeError):
    """材料里没有可识别的人声（纯音乐 / 纯画面 / 静音）：区别于基础设施故障，提示文案面向用户。"""


def _asr_sentences(file_path: str) -> list[tuple[str, str]]:
    """paraformer-v2 文件转写，返回 [(mm:ss, 句子文本)]"""
    oss_url = _upload_to_oss(file_path)
    res = requests.post(
        f'{ASR_BASE}/api/v1/services/audio/asr/transcription',
        headers={
            'Authorization': f'Bearer {settings.DASHSCOPE_API_KEY}',
            'X-DashScope-OssResourceResolve': 'enable',
            'X-DashScope-Async': 'enable',
        },
        json={'model': ASR_MODEL,
              'input': {'file_urls': [oss_url]},
              'parameters': {'language_hints': ['zh', 'en']}},
        timeout=60,
    )
    res.raise_for_status()
    task_id = res.json()['output']['task_id']
    for _ in range(ASR_POLL_MAX):
        time.sleep(ASR_POLL_INTERVAL)
        out = requests.get(f'{ASR_BASE}/api/v1/tasks/{task_id}',
                           headers={'Authorization': f'Bearer {settings.DASHSCOPE_API_KEY}'},
                           timeout=30).json().get('output', {})
        status_ = out.get('task_status')
        if status_ == 'SUCCEEDED':
            break
        if status_ == 'FAILED':
            message = out.get('message') or 'ASR task failed'
            # 无有效语音片段时任务会以 FAILED 落地（实测纯音乐/静音视频），单独识别给友好文案
            if 'SUCCESS_WITH_NO_VALID_FRAGMENT' in message:
                raise NoSpeechError(NO_SPEECH_TEXT)
            raise RuntimeError(f'转写任务失败：{message}')
    else:
        raise TimeoutError('转写任务超时')
    result = out['results'][0]
    url = result.get('transcription_url') or (result.get('output') or {}).get('file_url')
    data = requests.get(url, timeout=60).json()
    sentences = (data.get('transcripts') or [{}])[0].get('sentences') or []
    return [(_mmss(s.get('begin_time', 0)), s.get('text', '').strip()) for s in sentences if s.get('text', '').strip()]


def _vision_payload(file_path: str) -> tuple[str, bytes]:
    """图片转 base64 前先压缩：qwen-vl 单图（含 base64）上限 10MB，
    手机拍摄的板书照片常超限；等比缩到长边 2400px、JPEG q88 后一般 <3MB，识别效果不受影响"""
    mime = mimetypes.guess_type(file_path)[0] or 'image/jpeg'
    with open(file_path, 'rb') as f:
        raw = f.read()
    if len(raw) <= 6 * 1024 * 1024 and mime in ('image/jpeg', 'image/png', 'image/webp'):
        return mime, raw
    try:
        from io import BytesIO
        from PIL import Image
        img = Image.open(file_path)
        img = img.convert('RGB') if img.mode not in ('RGB', 'L') else img
        img.thumbnail((2400, 2400))
        buf = BytesIO()
        img.save(buf, format='JPEG', quality=88)
        data = buf.getvalue()
        if len(data) < len(raw):   # 压缩有效才采用
            return 'image/jpeg', data
    except Exception:
        pass  # 压缩失败退回原图，交由模型侧报错
    return mime, raw


def _vision_extract(file_path: str) -> str:
    """qwen-vl-plus 识别板书 / PPT 图片，转写出文字内容"""
    mime, raw = _vision_payload(file_path)
    b64 = base64.b64encode(raw).decode()
    text = llm.chat([{
        'role': 'user',
        'content': [
            {'type': 'image_url', 'image_url': {'url': f'data:{mime};base64,{b64}'}},
            {'type': 'text', 'text': '这是一张课堂板书或 PPT 照片。请完整转写图中的文字内容（含标题、公式、批注），'
                                     '按内容分条输出纯文本，不要添加图中没有的解释。'},
        ],
    }], model=llm.MODEL_VL, temperature=0.1)
    if not text:
        raise RuntimeError('图片识别失败')
    return text.strip()


def _read_text(file_path: str) -> str:
    with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
        return f.read(TEXT_MAX).strip()


def _pdf_text(file_path: str) -> str:
    """PyMuPDF 按页抽取文本；扫描件（无文本层）明确报错提示改传图片"""
    try:
        import pymupdf
    except ImportError:
        import fitz as pymupdf  # 旧版包名兼容
    pages = []
    with pymupdf.open(file_path) as doc:
        for page in doc:
            t = page.get_text().strip()
            if t:
                pages.append(t)
    text = '\n'.join(pages)
    if not text.strip():
        raise RuntimeError('PDF 中未提取到文本（可能是扫描件），请截图后作为图片材料上传')
    return text


def _docx_text(file_path: str) -> str:
    from docx import Document
    doc = Document(file_path)
    parts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(' | '.join(cells))
    if not parts:
        raise RuntimeError('Word 文档中未提取到文本')
    return '\n'.join(parts)


def _pptx_text(file_path: str) -> str:
    from pptx import Presentation
    prs = Presentation(file_path)
    parts = []
    for no, slide in enumerate(prs.slides, 1):
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                t = shape.text_frame.text.strip()
                if t:
                    texts.append(t)
            if getattr(shape, 'has_table', False) and shape.has_table:
                for row in shape.table.rows:
                    cells = [c.text.strip() for c in row.cells if c.text.strip()]
                    if cells:
                        texts.append(' | '.join(cells))
        if texts:
            parts.append(f'【第 {no} 页】\n' + '\n'.join(texts))
    if not parts:
        raise RuntimeError('PPT 中未提取到文本')
    return '\n'.join(parts)


def _xlsx_text(file_path: str) -> str:
    from openpyxl import load_workbook
    wb = load_workbook(file_path, read_only=True, data_only=True)
    parts = []
    for ws in wb.worksheets:
        rows = []
        for row in ws.iter_rows(values_only=True):
            cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
            if cells:
                rows.append(' | '.join(cells))
            if sum(len(r) for r in rows) > TEXT_MAX:
                break
        if rows:
            parts.append(f'【工作表 {ws.title}】\n' + '\n'.join(rows))
    wb.close()
    if not parts:
        raise RuntimeError('表格中未提取到内容')
    return '\n'.join(parts)[:TEXT_MAX]


def _csv_text(file_path: str) -> str:
    import csv as _csv
    rows = []
    with open(file_path, 'r', encoding='utf-8-sig', errors='ignore', newline='') as f:
        for row in _csv.reader(f):
            cells = [c.strip() for c in row if c.strip()]
            if cells:
                rows.append(' | '.join(cells))
            if sum(len(r) for r in rows) > TEXT_MAX:
                break
    if not rows:
        raise RuntimeError('CSV 中未提取到内容')
    return '\n'.join(rows)[:TEXT_MAX]


def _zip_member_name(info) -> str:
    """压缩包内文件名修复：Windows 资源管理器打 zip 不设 UTF-8 标志位，
    中文文件名在 zipfile 里默认按 cp437 解出乱码，需转回 GBK"""
    name = info.filename
    if not (info.flag_bits & 0x800):
        try:
            name = name.encode('cp437').decode('gbk', errors='ignore')
        except Exception:
            pass
    return name.replace('\\', '/')


def _read_zip(file_path: str) -> str:
    """压缩包抽取：逐个解出内嵌文档，按扩展名复用 _read_doc 解析后拼接。
    跳过目录 / 隐藏文件 / __MACOSX / 嵌套压缩包；单文件与总量设上限防压缩炸弹；
    单个内嵌文件解析失败时跳过（其余文件仍可生成），全部不可解析才报错。"""
    import tempfile
    import zipfile

    ZIP_DOC_EXTS = ('txt', 'md', 'markdown', 'pdf', 'docx', 'pptx', 'ppsx', 'xlsx', 'csv')
    parts: list[str] = []
    total = 0
    with zipfile.ZipFile(file_path) as zf:
        for info in zf.infolist():
            name = _zip_member_name(info)
            short = name.rsplit('/', 1)[-1]
            if info.is_dir() or name.startswith('.') or '__MACOSX' in name:
                continue
            ext = short.rsplit('.', 1)[-1].lower() if '.' in short else ''
            if ext == 'zip' or ext not in ZIP_DOC_EXTS:   # 嵌套压缩包不递归
                continue
            if info.file_size > ZIP_INNER_MAX:
                raise RuntimeError(f'压缩包内「{short}」超过 {ZIP_INNER_MAX // 1024 // 1024}MB，请拆分后再上传')
            total += info.file_size
            if total > ZIP_TOTAL_MAX:
                raise RuntimeError('压缩包解压后总内容超过 200MB，请拆分后再上传')
            data = zf.read(info)
            tmp_path = ''
            try:
                with tempfile.NamedTemporaryFile(suffix='.' + ext, delete=False) as tmp:
                    tmp.write(data)
                    tmp_path = tmp.name
                text = _read_doc(tmp_path)
            except Exception:
                continue   # 单文件解析失败不拖垮整包（如扫描件 PDF），临时文件由 finally 清理
            finally:
                if tmp_path and os.path.exists(tmp_path):
                    try:
                        os.unlink(tmp_path)
                    except OSError:
                        pass
            text = (text or '').strip()
            if text:
                parts.append(f'【文件：{short}】\n{text}')
    if not parts:
        raise RuntimeError('压缩包中没有可解析的文档（支持 txt / md / pdf / docx / pptx / xlsx / csv）')
    return '\n\n'.join(parts)[:TEXT_MAX]


def _read_doc(file_path: str) -> str:
    """讲义文档 / PPT / 压缩包抽取：按扩展名分派到对应解析器"""
    ext = os.path.splitext(file_path)[1].lower()
    if ext == '.zip':
        return _read_zip(file_path)
    if ext in ('.txt', '.md', '.markdown'):
        return _read_text(file_path)
    if ext == '.pdf':
        return _pdf_text(file_path)[:TEXT_MAX]
    if ext == '.docx':
        return _docx_text(file_path)[:TEXT_MAX]
    if ext in ('.pptx', '.ppsx'):
        return _pptx_text(file_path)[:TEXT_MAX]
    if ext == '.xlsx':
        return _xlsx_text(file_path)
    if ext == '.csv':
        return _csv_text(file_path)
    raise UnsupportedSource('该文档格式暂不支持自动生成纪要')


def _extract(note: Note) -> tuple[str, str]:
    """按材料类型抽取内容，返回 (内容文本, 来源标签)"""
    if not note.source:
        raise UnsupportedSource('未上传材料文件')
    with local_material_path(note.source) as path:
        return _extract_path(note, path)


def _extract_path(note: Note, path: str) -> tuple[str, str]:
    kind = note.source_type or 'audio'
    if kind in ('audio', 'video'):
        sentences = _asr_sentences(path)
        if not sentences:
            raise NoSpeechError(NO_SPEECH_TEXT)
        return '\n'.join(f'[{ts}] {text}' for ts, text in sentences), '录音'
    if kind == 'image':
        return _vision_extract(path), '板书'
    if kind == 'doc':
        return _read_doc(path), '讲义'
    if kind == 'ppt':
        return _read_doc(path), 'PPT'
    raise UnsupportedSource('该格式暂不支持自动生成纪要，请上传 音频 / 视频 / 图片 / PPT / PDF / Word / 文本 / 压缩包')


def _summarize(content: str, source_label: str) -> dict:
    """qwen-plus 融合生成纪要板块；时间戳（录音材料）对齐进 ref，供引用溯源"""
    system = (
        '你是课堂纪要编辑。根据给定的课堂材料内容生成结构化纪要，只输出 JSON：'
        '{"theme": "一句话课程主题",'
        ' "points": [{"text": "核心知识点", "ts": "mm:ss"}],'
        ' "definitions": [{"term": "术语", "text": "解释", "ts": "mm:ss"}],'
        ' "emphasized": ["老师强调或布置的重点问题"]}'
        '要求：忠实于材料内容，不要编造材料里没有的信息；'
        'points 4~8 条、definitions 2~6 条、emphasized 1~4 条；'
        '材料行首带 [mm:ss] 时间戳时，ts 必须取材料中真实存在的时间点；没有时间戳则省略 ts。'
    )
    data = llm.chat_json(
        [{'role': 'system', 'content': system},
         {'role': 'user', 'content': f'课堂材料（来源：{source_label}）：\n{content}'}],
        temperature=0.3,
        timeout=90,   # 长材料 JSON 输出可能超默认 30s，超时即整篇纪要静默失败
    )
    if not isinstance(data, dict):
        raise RuntimeError('纪要生成失败（模型输出异常）')

    stamps = set(re.findall(r'\[(\d{1,2}:\d{2})\]', content))
    has_ts = source_label == '录音' and bool(stamps)

    def _ref(ts) -> str:
        return f'录音 {ts}' if has_ts and ts in stamps else source_label

    points = []
    for p in (data.get('points') or []):
        if isinstance(p, dict) and p.get('text'):
            points.append({'text': str(p['text']).strip(), 'ref': _ref(p.get('ts'))})
        elif isinstance(p, str) and p.strip():
            points.append({'text': p.strip(), 'ref': source_label})
    definitions = []
    for d in (data.get('definitions') or []):
        if isinstance(d, dict) and d.get('term') and d.get('text'):
            definitions.append({'term': str(d['term']).strip(), 'text': str(d['text']).strip(), 'ref': _ref(d.get('ts'))})
    emphasized = [str(x).strip() for x in (data.get('emphasized') or []) if str(x).strip()]
    return {
        'theme': str(data.get('theme') or '').strip(),
        'points': points,
        'definitions': definitions,
        'emphasized': emphasized,
    }


def _split_chunks(content: str, size: int = SUMMARIZE_CHUNK) -> list[str]:
    """按段落（行）边界把长文切成 ≤size 字符的片段；单行超长时兜底硬切"""
    chunks: list[str] = []
    buf: list[str] = []
    buf_len = 0
    for line in content.split('\n'):
        line = line.rstrip()
        if buf_len + len(line) + 1 > size and buf:
            chunks.append('\n'.join(buf))
            buf, buf_len = [], 0
        while len(line) > size:      # 单行超长（如压缩成一行的讲义）
            chunks.append(line[:size])
            line = line[size:]
        buf.append(line)
        buf_len += len(line) + 1
    if buf:
        chunks.append('\n'.join(buf))
    return [c for c in chunks if c.strip()]


def _summarize_part(content: str, source_label: str, idx: int, total: int) -> dict:
    """map 步骤：对单个片段抽取结构化要点（同一 JSON schema，不带跨段汇总职责）"""
    system = (
        '你是课堂纪要编辑。下面是长课堂材料的片段（共 %d 段，当前第 %d 段）。'
        '只根据该片段内容抽取结构化要点，只输出 JSON：'
        '{"theme": "一句话主题",'
        ' "points": [{"text": "核心知识点", "ts": "mm:ss"}],'
        ' "definitions": [{"term": "术语", "text": "解释", "ts": "mm:ss"}],'
        ' "emphasized": ["老师强调或布置的重点问题"]}'
        '要求：忠实于片段内容、不要编造；points 2~8 条、definitions 0~6 条、emphasized 0~4 条；'
        '行首带 [mm:ss] 时间戳时 ts 取真实时间点，否则省略 ts。' % (total, idx)
    )
    data = llm.chat_json(
        [{'role': 'system', 'content': system},
         {'role': 'user', 'content': f'课堂材料片段（来源：{source_label}）：\n{content}'}],
        temperature=0.2,
        timeout=90,
    )
    return data if isinstance(data, dict) else {}


def _merge_partials(partials: list[dict]) -> dict:
    """合并各片段结果：按内容去重、限量，作为 reduce 兜底也可直接使用"""
    def dedupe(items, cap):
        seen, out = set(), []
        for it in items:
            t = (it.get('text') if isinstance(it, dict) else str(it)).strip()
            if not t or t in seen:
                continue
            seen.add(t)
            out.append(it)
            if len(out) >= cap:
                break
        return out

    points, definitions, emphasized = [], [], []
    for p in partials:
        points.extend(x for x in (p.get('points') or []) if isinstance(x, dict) and x.get('text'))
        definitions.extend(x for x in (p.get('definitions') or []) if isinstance(x, dict) and x.get('term'))
        emphasized.extend(x for x in (p.get('emphasized') or []) if isinstance(x, str) and x.strip())
    return {
        'theme': next((str(p.get('theme') or '').strip() for p in partials if p.get('theme')), ''),
        'points': dedupe(points, 20),
        'definitions': dedupe(definitions, 12),
        'emphasized': dedupe(emphasized, 8),
    }


def _summarize_long(content: str, source_label: str) -> dict:
    """长文 map-reduce：分段抽取 → 合并去重 → 融合精简成最终 schema"""
    chunks = _split_chunks(content)
    partials = []
    for i, chunk in enumerate(chunks, 1):
        data = _summarize_part(chunk, source_label, i, len(chunks))
        if data:
            partials.append(data)
    merged = _merge_partials(partials)

    def render_points(items):
        return '\n'.join(f'- {x["text"]}' + (f' [{x["ts"]}]' if x.get('ts') else '') for x in items)

    candidate = (
        f'主题：{merged["theme"]}\n'
        f'知识点候选：\n{render_points(merged["points"]) or "（无）"}\n'
        f'术语候选：\n' + '\n'.join(f'- {x["term"]}：{x["text"]}' for x in merged['definitions'])
        + '\n重点问题候选：\n' + '\n'.join(f'- {x}' for x in merged['emphasized'])
    )
    system = (
        '你是课堂纪要主编。以下是对同一份课堂材料分段抽取后合并的候选要点，'
        '请融合成最终结构化纪要：合并重复与相近表述、按重要度排序，只输出 JSON：'
        '{"theme": "一句话课程主题",'
        ' "points": [{"text": "核心知识点", "ts": "mm:ss"}],'
        ' "definitions": [{"term": "术语", "text": "解释", "ts": "mm:ss"}],'
        ' "emphasized": ["老师强调或布置的重点问题"]}'
        '要求：忠实于候选要点、不编造；points 4~8 条、definitions 2~6 条、emphasized 1~4 条；'
        '候选中带 [mm:ss] 时保留该 ts，没有则省略。'
    )
    data = llm.chat_json(
        [{'role': 'system', 'content': system},
         {'role': 'user', 'content': candidate}],
        temperature=0.2,
        timeout=90,
    )
    if isinstance(data, dict) and data.get('points'):
        return data
    return merged   # reduce 失败时兜底用合并结果（已去重限量）


def process_note(note_id: int) -> None:
    """流水线主体（后台线程执行）：抽取 → 生成（长文走分段 map-reduce）→ 入库；失败置为 failed"""
    note = Note.objects.filter(pk=note_id).first()
    if note is None:
        return
    try:
        content, source_label = _extract(note)
        fields = (_summarize(content, source_label) if len(content) <= SUMMARIZE_DIRECT_MAX
                  else _summarize_long(content, source_label))
        note.theme = fields['theme'] or note.title
        note.points = fields['points']
        note.definitions = fields['definitions']
        note.emphasized = fields['emphasized']
        note.transcript = content
        if source_label == '录音':
            stamps = re.findall(r'^\[(\d{1,2}:\d{2})\]', content, flags=re.M)
            note.duration = stamps[-1] if stamps else note.duration  # 末句时间 ≈ 时长
        note.status = Note.Status.DONE
        note.fail_reason = ''
        note.save()
        rag.index_note(note)  # 生成完成即向量化入库，供智能问答检索
    except NoSpeechError as exc:
        # 无可用语音属正常业务失败：原因直接面向用户展示，不打错误日志
        log.info('纪要无有效语音 note=%s (%s)', note_id, note.source.path if note.source else '?')
        note.status = Note.Status.FAILED
        note.fail_reason = str(exc)
        note.save(update_fields=['status', 'fail_reason', 'updated_at'])
    except Exception as exc:
        log.exception('纪要流水线失败 note=%s', note_id)
        note.status = Note.Status.FAILED
        note.fail_reason = str(exc)[:200] or '材料处理异常，请重试或更换材料'
        note.save(update_fields=['status', 'fail_reason', 'updated_at'])


def start_processing(note: Note) -> None:
    """置为处理中并启动后台流水线线程（接口立即返回，前端轮询状态）"""
    note.status = Note.Status.PROCESSING
    note.fail_reason = ''  # 重新处理时清掉上次的失败原因，避免处理中仍显示旧文案
    note.save(update_fields=['status', 'fail_reason', 'updated_at'])
    threading.Thread(target=process_note, args=(note.pk,), daemon=True).start()
