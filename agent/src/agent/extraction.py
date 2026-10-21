"""File extraction for registered, host-authorized classroom assets."""

import io
import os
from pathlib import Path
import re
import shutil
import subprocess
import tempfile
import zipfile

from .materials import Segment, timestamp


class ExtractionError(ValueError):
    pass


class Extractor:
    def __init__(self, ocr=None, max_bytes=50 * 1024 * 1024, max_pages=300):
        self.ocr = ocr
        self.max_bytes = max_bytes
        self.max_pages = max_pages

    def _ocr(self, data, ref):
        if self.ocr is None:
            raise ExtractionError(f'{ref} 需要 OCR；请配置视觉模型，未写入不完整材料')
        from PIL import Image
        with Image.open(io.BytesIO(data)) as image:
            image.thumbnail((2400, 2400))
            buffer = io.BytesIO()
            image.convert('RGB').save(buffer, format='PNG')
        text = self.ocr(buffer.getvalue(), 'image/png')
        if not isinstance(text, str) or not text.strip() or text.strip() == '[无法辨认]':
            raise ExtractionError(f'{ref} OCR 无可用结果，未写入材料')
        return Segment(text=text.strip(), ref=ref + '（OCR）')

    def extract(self, path):
        path = Path(path)
        if not path.is_file() or not 0 < path.stat().st_size <= self.max_bytes:
            raise ExtractionError('文件不存在、为空或超过 50MB 限制')
        suffix = path.suffix.lower()
        if suffix in ('.doc', '.ppt'):
            return self._legacy(path)
        if suffix in ('.docx', '.pptx'):
            with zipfile.ZipFile(path) as archive:
                if sum(item.file_size for item in archive.infolist()) > 200 * 1024 * 1024:
                    raise ExtractionError('Office 解压后大小超过限制')
        handlers = {'.txt': self._text, '.md': self._text, '.docx': self._docx,
                    '.pdf': self._pdf, '.pptx': self._pptx}
        if suffix in ('.png', '.jpg', '.jpeg', '.webp'):
            segments = [self._ocr(path.read_bytes(), path.name)]
        elif suffix in handlers:
            segments = handlers[suffix](path)
        else:
            raise ExtractionError('暂不支持此格式；录音请注册可供 ASR 读取的 HTTPS URL')
        if not segments:
            raise ExtractionError('文件未抽取到可用文字，未写入材料')
        if sum(len(s.text) for s in segments) > 2_000_000:
            raise ExtractionError('抽取文字超过 200 万字限制')
        return segments

    def _text(self, path):
        raw = path.read_bytes()
        try:
            text = raw.decode('utf-8-sig')
        except UnicodeDecodeError:
            text = raw.decode('gb18030')
        segments = []
        for line_no, line in enumerate(text.splitlines(), 1):
            line = line.strip()
            if not line:
                continue
            match = re.match(r'^\[(\d{2,}):([0-5]\d)\]\s*(.*)', line)
            if match and match[3].strip():
                ms = (int(match[1]) * 60 + int(match[2])) * 1000
                segments.append(Segment(text=match[3], ref=f'{path.name} 录音 {timestamp(ms)}', timestamp_ms=ms))
            else:
                segments.append(Segment(text=line, ref=f'{path.name} 第{line_no}行'))
        return segments

    def _docx(self, path):
        from docx import Document
        from docx.table import Table
        document = Document(path)
        segments = []
        for index, block in enumerate(document.iter_inner_content(), 1):
            if isinstance(block, Table):
                text = '\n'.join(' | '.join(cell.text for cell in row.cells) for row in block.rows)
            else:
                text = block.text
            if text.strip():
                segments.append(Segment(text=text.strip(), ref=f'{path.name} 第{index}块'))
        for index, relation in enumerate(document.part.rels.values(), 1):
            if relation.reltype.endswith('/image') and not relation.is_external:
                segments.append(self._ocr(relation.target_part.blob, f'{path.name} 图片{index}'))
        return segments

    def _pdf(self, path):
        from pypdf import PdfReader
        reader = PdfReader(path)
        if reader.is_encrypted:
            raise ExtractionError('请先提供已解密的 PDF')
        if len(reader.pages) > self.max_pages:
            raise ExtractionError('PDF 页数超过限制')
        segments = []
        pdf = None
        try:
            for index, page in enumerate(reader.pages):
                ref = f'{path.name} 第{index + 1}页'
                text = (page.extract_text() or '').strip()
                if text:
                    segments.append(Segment(text=text, ref=ref))
                else:
                    if self.ocr is None:
                        raise ExtractionError(f'{ref} 无文字层，需要 OCR')
                    import pypdfium2
                    if pdf is None:
                        pdf = pypdfium2.PdfDocument(str(path))
                    rendered_page = pdf[index]
                    bitmap = None
                    try:
                        width, height = rendered_page.get_size()
                        scale = min(2., 2400 / max(width, height))
                        bitmap = rendered_page.render(scale=scale)
                        image = bitmap.to_pil()
                        data = io.BytesIO()
                        image.save(data, format='PNG')
                        segments.append(self._ocr(data.getvalue(), ref))
                    finally:
                        if bitmap is not None:
                            bitmap.close()
                        rendered_page.close()
        finally:
            if pdf is not None:
                pdf.close()
        return segments

    def _pptx(self, path):
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE
        deck = Presentation(path)
        if len(deck.slides) > self.max_pages:
            raise ExtractionError('PPT 页数超过限制')
        segments = []
        def shapes(items, ref):
            for shape in items:
                if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                    shapes(shape.shapes, ref)
                elif shape.has_text_frame and shape.text.strip():
                    segments.append(Segment(text=shape.text.strip(), ref=ref))
                elif shape.has_table:
                    text = '\n'.join(' | '.join(cell.text for cell in row.cells) for row in shape.table.rows)
                    if text.strip():
                        segments.append(Segment(text=text.strip(), ref=ref))
                elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    segments.append(self._ocr(shape.image.blob, ref + ' 图片'))
        for index, slide in enumerate(deck.slides, 1):
            ref = f'{path.name} 第{index}页'
            shapes(slide.shapes, ref)
            if slide.has_notes_slide:
                frame = slide.notes_slide.notes_text_frame
                if frame is not None and frame.text.strip():
                    segments.append(Segment(text=frame.text.strip(), ref=ref + ' 备注'))
        return segments

    def _legacy(self, path):
        executable = shutil.which('soffice') or shutil.which('libreoffice')
        if not executable:
            raise ExtractionError('旧版 DOC/PPT 需要安装 LibreOffice，或另存为 DOCX/PPTX 后导入')
        extension = 'docx' if path.suffix.lower() == '.doc' else 'pptx'
        with tempfile.TemporaryDirectory(prefix='agent-convert-') as folder:
            profile = Path(folder, 'profile').as_uri()
            command = [executable, f'-env:UserInstallation={profile}', '--headless', '--convert-to', extension, '--outdir', folder, str(path.resolve())]
            result = subprocess.run(command, capture_output=True, timeout=120,
                                    creationflags=subprocess.CREATE_NO_WINDOW if os.name == 'nt' else 0)
            converted = Path(folder, path.stem + '.' + extension)
            if result.returncode or not converted.is_file():
                raise ExtractionError('LibreOffice 转换失败，请另存为新格式后重试')
            return [part.model_copy(update={'ref': part.ref.replace(converted.name, path.name, 1)}) for part in self.extract(converted)]
