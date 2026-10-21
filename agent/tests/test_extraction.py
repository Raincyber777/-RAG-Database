import io
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document
from pptx import Presentation
from pptx.util import Inches
from PIL import Image
from pypdf import PdfWriter
from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject

from agent.extraction import Extractor, ExtractionError
from agent.materials import Segment, chunk_segments


class ExtractionTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)

    def test_txt_preserves_real_timestamps_and_source(self):
        path = self.root / '课堂.txt'
        path.write_text('[08:15] 导数表示变化率。\n下一段原文。', encoding='utf-8')
        result = Extractor().extract(path)
        self.assertEqual(result[0].timestamp_ms, 495000)
        self.assertIn('08:15', result[0].ref)
        self.assertIsNone(result[1].timestamp_ms)

    def test_docx_extracts_paragraph_and_table_in_order(self):
        doc = Document()
        doc.add_paragraph('第一段')
        table = doc.add_table(rows=1, cols=2)
        table.cell(0, 0).text = '术语'
        table.cell(0, 1).text = '解释'
        doc.add_paragraph('最后一段')
        path = self.root / '讲义.docx'
        doc.save(path)
        text = '\n'.join(s.text for s in Extractor().extract(path))
        self.assertLess(text.index('第一段'), text.index('术语'))
        self.assertLess(text.index('解释'), text.index('最后一段'))

    def test_pptx_text_tables_notes_and_image_ocr_have_slide_refs(self):
        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1)).text = '微积分'
        table = slide.shapes.add_table(1, 2, Inches(1), Inches(2), Inches(5), Inches(1)).table
        table.cell(0, 0).text = '导数'
        table.cell(0, 1).text = '变化率'
        slide.notes_slide.notes_text_frame.text = '讲解备注'
        image = io.BytesIO()
        Image.new('RGB', (20, 20), 'white').save(image, format='PNG')
        image.seek(0)
        slide.shapes.add_picture(image, Inches(1), Inches(3))
        path = self.root / '课件.pptx'
        deck.save(path)
        parts = Extractor(ocr=lambda data, mime: '图片中的公式').extract(path)
        self.assertIn('图片中的公式', '\n'.join(p.text for p in parts))
        self.assertIn('讲解备注', '\n'.join(p.text for p in parts))
        self.assertTrue(all('第1页' in p.ref for p in parts))

    def test_pdf_text_layer_and_scanned_page_ocr(self):
        path = self.root / 'lecture.pdf'
        writer = PdfWriter()
        page = writer.add_blank_page(width=300, height=300)
        font = DictionaryObject({NameObject('/Type'): NameObject('/Font'), NameObject('/Subtype'): NameObject('/Type1'), NameObject('/BaseFont'): NameObject('/Helvetica')})
        page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'): DictionaryObject({NameObject('/F1'): writer._add_object(font)})})
        stream = DecodedStreamObject()
        stream.set_data(b'BT /F1 12 Tf 10 200 Td (Derivative is a rate.) Tj ET')
        page[NameObject('/Contents')] = writer._add_object(stream)
        writer.add_blank_page(width=300, height=300)
        writer.write(path)
        calls = []
        def ocr(data, mime):
            calls.append(data)
            return '扫描页识别内容'
        result = Extractor(ocr=ocr).extract(path)
        self.assertIn('Derivative', result[0].text)
        self.assertIn('第2页', result[1].ref)
        self.assertEqual(len(calls), 1)
        self.assertTrue(calls[0].startswith(b'\x89PNG'))

    def test_scanned_pdf_without_ocr_does_not_silently_index_empty_text(self):
        path = self.root / 'scan.pdf'
        writer = PdfWriter()
        writer.add_blank_page(width=100, height=100)
        writer.write(path)
        with self.assertRaises(ExtractionError):
            Extractor().extract(path)

    def test_legacy_office_requires_converter(self):
        path = self.root / 'old.ppt'
        path.write_bytes(b'legacy')
        with patch('agent.extraction.shutil.which', return_value=None):
            with self.assertRaisesRegex(ExtractionError, 'LibreOffice'):
                Extractor().extract(path)

    def test_chunking_retains_location_and_covers_original_tail(self):
        parts = chunk_segments([Segment(text='知识点' * 1000 + '结尾', ref='讲义 第3页')])
        self.assertTrue(all(len(p.text) <= 800 and p.ref == '讲义 第3页' for p in parts))
        self.assertTrue(parts[-1].text.endswith('结尾'))
